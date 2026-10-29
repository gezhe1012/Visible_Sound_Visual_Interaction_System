import sys
from pptx import Presentation

src = sys.argv[1]
out = sys.argv[2]

prs = Presentation(src)
lines = []
for i, slide in enumerate(prs.slides, 1):
    lines.append(f"=== SLIDE {i} ===")
    for sh in slide.shapes:
        kind = sh.shape_type
        name = sh.name
        txt = ""
        if sh.has_text_frame:
            txt = " / ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
        lines.append(f"  - {name} ({kind}): {txt}")
        if sh.has_table:
            for r in sh.table.rows:
                lines.append("     | " + " || ".join(c.text.replace('\n',' ') for c in r.cells))
    if slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame.text.strip()
        if nt:
            lines.append("  NOTES: " + nt.replace("\n", " / "))

with open(out, 'w', encoding='utf-8') as f:
    f.write("\n".join(lines))
print("slides:", len(prs.slides), "chars:", sum(len(l) for l in lines))
