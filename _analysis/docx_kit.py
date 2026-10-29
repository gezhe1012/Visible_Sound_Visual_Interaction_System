# -*- coding: utf-8 -*-
"""DOCX 构建辅助库：为修订版技术文档提供统一的排版原语。

设计基线（沿用原文档）：
- A4，上下 2.54 cm、左右 3.17 cm；
- 正文 宋体 / Times New Roman 12 pt，1.5 倍行距，首行缩进 2 字符；
- 部分标题 黑体 14 pt 加粗；一级标题 黑体 12 pt；二级标题 宋体加粗 12 pt；
- 图题、表题 9.5 pt，灰色 404040，居中；
- 表格 9.5 pt，表头加粗，单线边框。
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_FONT_CN = "宋体"
BODY_FONT_EN = "Times New Roman"
HEAD_FONT_CN = "黑体"

CAPTION_COLOR = RGBColor(0x40, 0x40, 0x40)
TEXT_WIDTH_CM = 21.0 - 2 * 3.17  # 14.66 cm 版心宽度


def _set_run_font(run, cn=BODY_FONT_CN, en=BODY_FONT_EN, size=12.0,
                  bold=None, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), en)
    rfonts.set(qn('w:hAnsi'), en)
    rfonts.set(qn('w:eastAsia'), cn)


def _set_para_format(p, first_indent_chars=0, left_cm=0.0, space_before=0.0,
                     space_after=0.0, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = line
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if left_cm:
        pf.left_indent = Cm(left_cm)
    if first_indent_chars:
        # 用字符数缩进，保证中英文混排时视觉为 2 字符
        ind = p._p.get_or_add_pPr().find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            p._p.get_or_add_pPr().append(ind)
        ind.set(qn('w:firstLineChars'), str(int(first_indent_chars * 100)))
        ind.set(qn('w:firstLine'), str(int(first_indent_chars * 240)))
    else:
        pf.first_line_indent = Pt(0)


def body(doc, text, indent=True, size=12.0, bold=False, align=None,
         space_after=0.0, left_cm=0.0):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=2 if indent else 0, left_cm=left_cm,
                     space_after=space_after,
                     align=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY)
    r = p.add_run(text)
    _set_run_font(r, size=size, bold=bold)
    return p


def rich(doc, segments, indent=True, size=12.0, align=None, space_after=0.0):
    """segments: [(text, bold), ...]"""
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=2 if indent else 0,
                     space_after=space_after,
                     align=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY)
    for text, bold in segments:
        r = p.add_run(text)
        _set_run_font(r, size=size, bold=bold)
    return p


def part(doc, text, first=False):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=0,
                     space_before=0 if first else 16, space_after=8,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _set_run_font(r, cn=HEAD_FONT_CN, en=HEAD_FONT_CN, size=14.0, bold=True)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=0, space_before=10, space_after=4,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _set_run_font(r, cn=HEAD_FONT_CN, en=HEAD_FONT_CN, size=12.0, bold=True)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=0, space_before=6, space_after=2,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _set_run_font(r, cn=BODY_FONT_CN, en=BODY_FONT_EN, size=12.0, bold=True)
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=0, space_before=4, space_after=2,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    _set_run_font(r, cn=BODY_FONT_CN, en=BODY_FONT_EN, size=12.0, bold=True)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=0, space_before=3, space_after=8,
                     align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    _set_run_font(r, size=9.5, color=CAPTION_COLOR)
    return p


def figure(doc, path, width_cm=13.5, caption_text=None):
    p = doc.add_paragraph()
    _set_para_format(p, first_indent_chars=0, space_before=6, space_after=2,
                     align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if caption_text:
        caption(doc, caption_text)


def _border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '808080')
        borders.append(el)
    tblPr.append(borders)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def table(doc, header, rows, widths_cm=None, size=9.5, caption_text=None,
          first_col_bold=False, header_fill='F2F2F2'):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _border(t)

    def _fill(cells, values, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
        for c, v in zip(cells, values):
            c.text = ""
            p = c.paragraphs[0]
            _set_para_format(p, first_indent_chars=0, line=1.0, space_before=1,
                             space_after=1, align=align)
            r = p.add_run(str(v))
            _set_run_font(r, size=size, bold=bold)

    _fill(t.rows[0].cells, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in t.rows[0].cells:
        _shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        _fill(cells, row)
        if first_col_bold:
            for r in cells[0].paragraphs[0].runs:
                r.font.bold = True

    if widths_cm:
        for row in t.rows:
            for c, w in zip(row.cells, widths_cm):
                c.width = Cm(w)
    # 表头跨页重复
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    trPr.append(th)
    if caption_text:
        caption(doc, caption_text)
    return t


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run = p.add_run()
    _set_run_font(run, size=9.0)
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
